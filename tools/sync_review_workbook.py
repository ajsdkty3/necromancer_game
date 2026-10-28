"""Synchronize the user-edited workbook in place, never regenerate from game data."""
from pathlib import Path
from shutil import copy2
from datetime import datetime
import re
from openpyxl import load_workbook
from openpyxl.comments import Comment

root = Path(__file__).resolve().parents[1]
path = root / '单位数据表.xlsx'
backup = root / 'backups'
backup.mkdir(exist_ok=True)
copy2(path, backup / f'单位数据表_同步前_{datetime.now():%Y%m%d_%H%M%S}.xlsx')
book = load_workbook(path)
allies, enemies = book['己方'], book['敌人']
original = {(s.title,c.coordinate): c.value for s in book for row in s for c in row}
by_name = {allies.cell(r,1).value:r for r in range(2,allies.max_row+1)}

# Repair a stale maximum in the healing description using the edited HP cell.
r = by_name['吸血鬼']
allies.cell(r,6).value = re.sub(r'不超过\d+生命上限',f'不超过{allies.cell(r,2).value}生命上限',allies.cell(r,6).value)
for sheet in [allies]:
    for r in range(2,sheet.max_row+1):
        for col in [5,6]:
            cell = sheet.cell(r,col)
            cell.value = cell.value.replace('相邻敌方单体','相邻指定单位单体（可指定友军）')
            note = '\n主动攻击均可指定友军；范围攻击覆盖的友军同样受到伤害。'
            cell.value += note
        name = sheet.cell(r,1).value
        if name == '亡灵弓兵':
            sheet.cell(r,6).comment = Comment('待确认：原文“基础伤害59点”是9点、5–9点还是59点？暂保留原文。','Codex')
        if name == '亡灵骑士':
            sheet.cell(r,6).comment = Comment('待确认：新践踏的8伤害、击退及跟进是否仍附带眩晕？原文眩晕说明暂保留。','Codex')

for r in range(2,enemies.max_row+1):
    source = by_name['亡灵'+enemies.cell(r,1).value]
    for col in [5,6]:
        src = allies.cell(source,col)
        enemies.cell(r,col).value = src.value.replace('支配力','AP')
        if src.comment:
            enemies.cell(r,col).comment = Comment(src.comment.text,src.comment.author)

damage = {'吸血鬼':'7–9','僵尸':'3–5','亡灵矛兵':'2','矛兵':'2','亡灵盾兵':'1','盾兵':'1','亡灵骑士':'5','骑士':'5'}
for sheet in book:
    for r in range(2,sheet.max_row+1):
        name = sheet.cell(r,1).value
        if '弓兵' in name:
            sheet.cell(r,7).value = '无（没有控制区，也没有借机攻击）'
        else:
            sheet.cell(r,7).value = (
                '控制区：1格。敌人从区域内继续移动或离开时，在离开旧格前触发；走入不触发。\n'
                '不消耗资源，每个单位仅可触发1次，在自身所属阵营回合开始时恢复；触发后红色范围消失。\n'
                f'仅触发移动的敌人受到{damage[name]}点单体基础伤害，不波及其他单位。'
                + ('骑士借机攻击基础伤害固定5点。' if '骑士' in name else '伤害数值与普通攻击一致。')
                + '\n借机攻击仅同步伤害数值，不附带普通攻击的范围伤害、击退、冲锋等效果。只覆盖和命中高度差≤1级的目标。'
                + ('\n众志成城照常增加伤害。' if sheet.title=='敌人' else '')
            )

book.save(path)
verified = load_workbook(path)
assert verified.sheetnames == ['己方','敌人']
for s in verified:
    for r in range(2,s.max_row+1):
        for col in [1,2,3,4,8]:
            c=s.cell(r,col)
            assert c.value == original[(s.title,c.coordinate)]
        assert '可指定友军' in s.cell(r,6).value
for r in range(2,enemies.max_row+1):
    source=by_name['亡灵'+enemies.cell(r,1).value]
    for col in [5,6]:
        assert verified['敌人'].cell(r,col).value == verified['己方'].cell(source,col).value.replace('支配力','AP')
print('Workbook synchronized; user HP/resources/movement/traits preserved; ambiguous entries commented.')
