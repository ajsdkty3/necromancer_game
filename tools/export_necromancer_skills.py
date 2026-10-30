from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

book = Workbook()
sheet = book.active
sheet.title = "亡灵法师技能"
sheet.append(["技能", "支配力消耗", "目标", "效果", "使用限制与说明"])
sheet.append(["召唤僵尸", 20, "场上任意未破坏的尸体", "摧毁尸体，在原格生成僵尸。最大生命10，当前生命继承尸体当前生命；例如尸体6/10→僵尸6/10。", "亡灵回合使用，次数不限，无距离限制。召唤后可按正常规则行动。"])
sheet.append(["血肉塑形", 20, "任意存活友方亡灵", "本场战斗最大生命值+10，同时恢复10生命；可重复施放并叠加。", "亡灵回合使用，次数不限，无距离限制，仍需足够支配力。"])
sheet.append(["献祭", 0, "任意存活友方亡灵", "直接杀死目标，按死亡前当前生命值等量恢复支配力（1血=1支配力）；魂归主宰独立+5。例如10血僵尸共恢复15。", "亡灵回合使用，次数不限，无距离限制。恢复受支配力上限限制。直接死亡，不被守护转移。亡灵不留下尸体。"])
for cell in sheet[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="354F46")
for row in sheet.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font = Font(name="Microsoft YaHei", size=11)
for col, width in {"A":18,"B":16,"C":28,"D":65,"E":65}.items():
    sheet.column_dimensions[col].width = width
for n in range(2,5):
    sheet.row_dimensions[n].height = 95
sheet.freeze_panes = "A2"
sheet.auto_filter.ref = sheet.dimensions
book.save(Path(__file__).resolve().parents[1] / "亡灵法师技能表.xlsx")
