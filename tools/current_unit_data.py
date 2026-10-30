"""Current data-derived action descriptions shared by documents and spreadsheets."""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT/'necromancer-test'/'data'
definitions = json.loads((DATA/'units.json').read_text(encoding='utf-8'))
members = json.loads((DATA/'battle_setup.json').read_text(encoding='utf-8'))['units']
ADVANTAGE = '须自身高于目标，或敌方目标相邻六格至少3名己方存活单位；尸体不计包围。仍须符合攻击范围与高度限制；可指定友军，但友军目标须高度优势。'
PUSH = '击退1格；后方障碍、单位（含尸体）、更高地形或越界阻止位移，伤害照常。下降另结算每级2点坠落伤害；击退不触发借机。'
def damage(d):
    return str(d['damage_min']) if d['damage_min']==d['damage_max'] else f"{d['damage_min']}–{d['damage_max']}"
def ordinary(d,r):
    if d['kind']=='BATTLE_ORDER':
        return f"战斗号令：{d['cost']}{r}，相邻6格友军伤害+1，不含自己；同类不叠加，持续至施放者下次己方回合开始。金黄色范围持续显示。"
    if d['kind']=='CHARGE':
        return f"冲锋：{d['cost']}{r}，沿六方向直线移动恰好{d['charge_range']}格，不另收移动费，移动不触发借机。终点仅正前方1格单体基础{damage(d)}伤害，含友伤；无目标仍可冲锋。单位、尸体、障碍或相邻高差>1阻挡路径；目标高差≤1。"+PUSH
    text=f"{d['label']}：{d['cost']}{r}，范围{d['range']}，基础伤害{damage(d)}；可指定友军。"
    if d['target_shape']=='spear':
        text+='直线目标覆盖该方向前两格（含友军）；非直线仅指定目标。中间障碍或高于攻击者的地形阻挡第二格；非直线距离2须至少一条最短路线通畅。单位、尸体可穿透。'
    if d['kind']=='SHOOT':
        text+='定点单体；中间单位、尸体不阻挡也不受伤。向上高差≤2，向下不限。'
    else:
        text+='近战目标高差≤1。'
    if d.get('ordinary_knockback'): text+=PUSH
    return text
def special(d,r):
    if 'stomp_cost' in d:
        text=f"战争践踏：{d['stomp_cost']}{r}，范围{d['stomp_range']}，基础伤害{d['stomp_damage_min']}；击退或击杀后，目标旧格为空时跟进。不附带眩晕。"+PUSH
    elif 'knockback_cost' in d:
        text=f"击退：{d['knockback_cost']}{r}，范围{d['knockback_range']}单体，基础伤害{d['knockback_damage']}。"+PUSH
    else:
        text=f"{d['advantage_label']}：{d.get('advantage_cost',d['cost'])}{r}，范围{d['range']}。"
        if 'advantage_hits' in d:
            text+=f"连续{d['advantage_hits']}次，每次独立{d['advantage_damage_min']}–{d['advantage_damage_max']}伤害；只扣一次费用，目标死亡即停止。"
        else: text+=f"基础伤害{d['overhead_damage']}。"
        if 'advantage_heal' in d: text+=f"恢复自己{d['advantage_heal']}血，不超过当前最大生命。"
        if d.get('advantage_stun'): text+='眩晕：目标下个所属阵营回合不能行动，重复不延长；不击退。'
        if d.get('advantage_vulnerable'): text+='先伤害后施加易损：所受全部伤害+50%，向上取整；持续至施放者下次己方回合开始，重复只刷新、不叠加；尸体免疫。'
    text+=ADVANTAGE
    if 'teleport_cost' in d: text+=f"\n瞬移：{d['teleport_cost']}{r}，{d['teleport_range']}格内空格；可越过路径阻挡，不得落在障碍或占用格，不触发借机。"
    if 'guard_cost' in d: text+=f"\n守护：{d['guard_cost']}{r}，相邻6格友军免伤并转移给自己，至下次己方回合开始。两名开启守护者互不保护；多人守护按部署顺序仅转移一次。范围伤害命中自身及友军时分别承担。白色范围持续显示。"
    return text
def reaction(d):
    if d['zoc']==0: return '无（无控制区）'
    return f"范围{d['zoc']}，单体基础伤害{d.get('reaction_damage',damage(d))}；每次己方回合开始恢复1次；离开旧格前结算，无资源消耗。仅攻击触发敌人，不穿刺、不击退、不附带优势动作效果。"
def rows(team):
    present={m['archetype'] for m in members if m['team']==team}
    for key,d in definitions.items():
        if key not in present: continue
        r='支配力' if team=='undead' else 'AP'
        yield [d['name'],d['role'],d['hp'],'共享支配力' if team=='undead' else f"{d.get('max_ap',6)} AP / 回合",f"同高每格{d['move_cost']}{r}；上下1级额外+1（当前2点/格）；高差>1不能步行。单位、尸体、障碍阻挡。",ordinary(d,r),special(d,r),reaction(d),'魂归主宰：死亡恢复5支配力，受上限限制；无士气，不留尸体。' if team=='undead' else '众志成城：每名相邻人类友军使自身伤害+1。初始士气100；死亡留下10血尸体。']
def export():
    book=Workbook(); book.remove(book.active)
    for team,title in [('undead','己方'),('human','敌人')]:
        s=book.create_sheet(title)
        s.append(['单位','类别','生命值（HP）','行动资源','移动','普通动作','优势动作 / 其他技能','借机攻击','特性'])
        for row in rows(team): s.append(row)
        s.freeze_panes='A2'; s.auto_filter.ref=s.dimensions
        for c in s[1]:
            c.font=Font(name='Microsoft YaHei',bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='354F46')
        for row in s.iter_rows(min_row=2):
            for c in row: c.font=Font(name='Microsoft YaHei',size=11); c.alignment=Alignment(wrap_text=True,vertical='top')
        for col,width in zip('ABCDEFGHI',[19,12,14,20,43,64,82,58,45]): s.column_dimensions[col].width=width
        for n in range(2,s.max_row+1): s.row_dimensions[n].height=240
    book.save(ROOT/'单位数据表.xlsx')
