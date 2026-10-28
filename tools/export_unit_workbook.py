"""Export a compact unit sheet with actions on each unit's own row."""
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "necromancer-test" / "data"
definitions = json.loads((DATA / "units.json").read_text(encoding="utf-8"))
members = json.loads((DATA / "battle_setup.json").read_text(encoding="utf-8"))["units"]
archetype_order = list(definitions)
members.sort(key=lambda member: archetype_order.index(member["archetype"]))
output = ROOT / "单位数据表.xlsx"
book = Workbook()
book.remove(book.active)
headers = ["单位", "生命值（HP）", "行动资源", "移动", "攻击", "优势动作 / 其他技能", "借机攻击", "特性"]

for team, title in [("undead", "己方"), ("human", "敌人")]:
    sheet = book.create_sheet(title)
    sheet.append(headers)
    seen = set()
    for member in members:
        if member["team"] != team or member["archetype"] in seen:
            continue
        seen.add(member["archetype"])
        d = definitions[member["archetype"]]
        resource = "支配力" if team == "undead" else "AP"
        damage = str(d["damage_min"]) if d["damage_min"] == d["damage_max"] else f'{d["damage_min"]}–{d["damage_max"]}'
        movement = f'同高每格{d["move_cost"]}{resource}；上升或下降1级额外消耗1{resource}（当前每格2点）。\n高度差超过1级不可步行跨越；可连续移动，不能穿越其他单位或障碍格。'
        attack = f'{d["label"]}：消耗{d["cost"]}{resource}，攻击{d["range"]}格范围内的指定目标。\n'
        if d["target_shape"] == "spear":
            attack += (
                f'目标在六条直线方向上：对该方向前1～2格内的所有单位造成范围伤害，每个目标独立受到{damage}伤害。\n'
                f'目标不在直线上：仅对指定目标造成{damage}伤害。\n可伤害友军。'
            )
        else:
            attack += f'造成{damage}单体伤害，可指定友军。'
            if d["kind"] == "SHOOT":
                attack += "\n中间单位不阻挡，也不会受伤。"
        if d["kind"] != "SHOOT":
            attack += "\n仅能命中与攻击者高度差≤1级的单位；范围攻击逐个目标检查高度。"
        if d["kind"] == "CHARGE":
            attack = f'骑士冲锋：整次消耗{d["cost"]}{resource}，直线移动恰好{d["charge_range"]}格，不另收移动费。冲锋移动不触发借机攻击。\n到达后自动攻击正前方、左前方、右前方3格中的所有单位，含友军；每个目标基础{damage}伤害，并向远离骑士方向击退1格。无目标则结束。\n不能穿过单位、障碍或相邻高差超过1级的地形；命中仍须高差≤1。击退后方为障碍、任意单位、更高地形或地图外时不位移，伤害仍结算。下降另结算每级2点坠落伤害，击退不触发借机。\n人类AI仅在命中敌人数多于友军数时冲锋，否则尝试换位寻找角度。'
        special = "无"
        if "teleport_range" in d:
            special = (
                f'瞬移：消耗{d["teleport_cost"]}{resource}，到达自身{d["teleport_range"]}格范围内任一空格。\n'
                '可越过单位和障碍，但不能落在障碍格；不触发借机攻击，之后仍可继续行动。'
            )
        if "guard_cost" in d:
            special = (
                f'守护：消耗{d["guard_cost"]}{resource}，持续至下次所属阵营回合开始。\n'
                '相邻6格友军所受伤害全部转移到自己；友军微颤后闪白光、不扣血。\n'
                '双方都开启守护的盾兵互不保护；只开一方可保护未开启的一方。\n'
                '同一友军被多人守护时按部署顺序转移一次。范围伤害命中自己及友军时，两份都承担。'
            )
        overhead = (
            f'优势动作·当头猛击：消耗{d["cost"]}{resource}，自身高于目标，或敌方目标相邻六格内有至少3名存活己方单位（包括行动者）；目标仍须在普通攻击范围内。'
            + (f'基础伤害{d["overhead_damage"]}。' if 'overhead_damage' in d else (f'基础伤害{d["damage_min"]*2}。' if d['damage_min'] == d['damage_max'] else f'基础伤害{d["damage_min"]*2}–{d["damage_max"]*2}。'))
            + ('每名相邻人类友军额外增加2伤害。' if team == 'human' else '')
            + '守护转移不重复加成。'
            + ('近战高度差不得超过1级。' if d['kind'] != 'SHOOT' else '远程无高度差上限。')
            + '沿用普通攻击形状和友伤规则，每个目标分别检查优势条件。'
        )
        if "knockback_cost" in d:
            overhead = '优势动作·击退：3AP，范围1格单体，基础伤害2；自身高于目标，或敌方目标相邻六格内有至少3名存活己方单位（包括行动者）。近战高度差不得超过1级。\n目标后退1格，后方有障碍、任意单位、更高地势或越界则不移动，伤害仍结算。\n落到较低地势时另受每级2点坠落伤害；坠落不叠加众志成城。击退不触发借机攻击。'
            overhead = overhead.replace('3AP',f'{d["knockback_cost"]}{resource}').replace('基础伤害2',f'基础伤害{d["knockback_damage"]}').replace('每级2点',f'每级{d.get("fall_damage_per_height",2)}点')
        if "advantage_label" in d:
            overhead = f'优势动作·{d["advantage_label"]}：消耗{d["cost"]}{resource}，范围{d["range"]}格。\n'
            overhead += '自身高于目标，或敌方目标相邻六格内有至少3名存活己方单位（包括行动者）。\n'
            if "advantage_hits" in d:
                overhead += f'连续造成{d["advantage_hits"]}次伤害，每次独立造成{d["advantage_damage_min"]}–{d["advantage_damage_max"]}点；目标死亡或战斗结束即停止。整次只扣一次消耗。\n'
            else:
                overhead += f'基础伤害{d["overhead_damage"]}点。'
            if "advantage_heal" in d:
                overhead += f'恢复自身{d["advantage_heal"]}生命，不超过{d["hp"]}生命上限。\n'
            overhead += '其余沿用普通攻击的目标和友伤规则。'
            overhead += '远程无高度差上限；中间单位不阻挡且不受伤。' if d['kind']=='SHOOT' else '近战高度差不得超过1级。'
            if team == 'human':
                overhead += '每名相邻人类友军额外增加1伤害；守护转移不重复加成。'
        if "stomp_cost" in d:
            overhead = f'优势动作·战争践踏：消耗{d["stomp_cost"]}{resource}，范围1，相邻敌方单体4–6伤害并眩晕。\n自身高于目标，或目标被至少3名己方单位包围；近战高差≤1。\n眩晕立即禁止行动及借机攻击，持续至目标下个所属阵营回合结束。重复眩晕不叠加或刷新；人类与亡灵均适用。'
        special = overhead if special == "无" else overhead + '\n\n' + special
        if d['kind'] == 'CHARGE':
            attack = attack.replace('人类AI仅在', '人类AI进攻时在')
            attack += '\n需要离开有效敌方控制区时，优先用冲锋脱离，可冲向无目标地点；仍避免友伤占优的路线。'
        reaction = "无（没有控制区）"
        if d["zoc"]:
            reaction = (
                f'控制区：{d["zoc"]}格。敌人从区域内继续移动或离开时，在离开旧格前触发；走入不触发。\n'
                f'不消耗资源，每个单位仅可触发1次，在自身所属阵营回合开始时恢复；触发后红色范围消失。仅触发移动的敌人受到{damage}单体伤害，不波及其他单位。'
            )
            if d["target_shape"] == "spear":
                reaction += "\n不使用普通穿刺的直线范围伤害；直线与非直线目标均只命中触发者。"
            reaction += "\n只覆盖和命中与攻击者高度差≤1级的格子及单位。"
        sheet.append([
            d["name"], d["hp"],
            "共享支配力" if team == "undead" else f'独立AP：每回合{d.get("max_ap",6)}点。\n' + "初始士气100；相邻六格敌人3–4名-20、5名-40、6名-80（可解除的状态惩罚）。2格内每名友军阵亡累计-20。HP小于等于20%时-20、初始人数阵亡过半时-20（状态惩罚，不重复叠加）。\n士气0时溃逃2个人类回合，只向随机可达右边缘移动，禁用攻击/守护/借机；到达右边缘离场，不算阵亡。士气0且六面被敌人包围时无法移动或恢复，解除包围后重新计算2个溃逃回合；再次完全被围则清空计时并保持士气0。未离场则基础士气恢复到80，再扣当前状态惩罚；仅低血量与阵亡过半生效时为40士气，包围惩罚另计。",
            movement, attack, special, reaction,
            ("精英。" if d["role"]=="精英" else "") + ("魂归主宰：死亡时恢复5点共享支配力，不超过本关支配力上限。" if team == "undead" else "众志成城：每有1名相邻六格内的存活人类友军，造成的伤害+1。普通攻击、借机攻击、范围攻击每个目标均适用；守护转移不重复加成。不再减免受到的伤害。"),
        ])
    table = Table(displayName="AlliedUnits" if team == "undead" else "EnemyUnits",
                  ref=f"A1:H{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 80
    sheet.freeze_panes = "D2"
    sheet.row_dimensions[1].height = 32
    for cell in sheet[1]:
        cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=12)
        cell.fill = PatternFill("solid", fgColor="243E50")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = 360 if team == 'human' else 300
        for cell in row:
            cell.font = Font(name="微软雅黑", size=11, color="243746")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[1].alignment = Alignment(horizontal="center", vertical="top")
    for index, width in enumerate([16, 17, 25, 33, 63, 44, 63, 44], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_area = sheet.dimensions
    sheet.print_title_rows = "1:1"

book.save(output)
checked = load_workbook(output)
assert checked.sheetnames == ["己方", "敌人"]
for sheet in checked:
    assert sheet.max_row == (7 if sheet.title == "己方" else 5) and sheet.max_column == 8
    assert [cell.value for cell in sheet[1]] == headers
assert all(checked["己方"].cell(row, 3).value == "共享支配力" for row in range(2, 8))
assert [checked["己方"].cell(row, 2).value for row in range(2, 4)] == [15, 8]
assert "消耗6支配力" in checked["己方"]["F2"].value
assert "范围伤害" in checked["敌人"]["E2"].value
assert checked["敌人"]["B4"].value == 20
assert "消耗3AP" in checked["敌人"]["F4"].value
assert all("命中率" not in str(cell.value) for sheet in checked for row in sheet for cell in row)
for sheet in checked:
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] in ['骑士', '亡灵骑士']:
            assert '恰好3格' in row[4] and '基础4伤害' in row[4]
            assert '整次消耗5' in row[4] and '冲锋移动不触发借机攻击' in row[4]
        if row[0] in ['弓兵', '亡灵弓兵']:
            assert '要害射击' in row[5] and '基础伤害5点' in row[5]
        if row[0] == '吸血鬼':
            assert '猩红之吻' in row[5] and '恢复自身4生命' in row[5]
        if row[0] == '僵尸':
            assert 'Feeding Frenzy' in row[5] and '连续造成3次' in row[5] and '1–2点' in row[5]
print("Workbook updated and verified: two sheets; unit actions inline; no stage, position or detail sections.")
